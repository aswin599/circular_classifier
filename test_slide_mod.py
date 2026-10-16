import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = pptx.Presentation('TCE_Linkedin_Project_Viva.pptx')
slide = prs.slides[0]

# Check Slide 1 shapes
for shape in slide.shapes:
    print(shape.name, shape.shape_type)

print("Slide width, height:", prs.slide_width, prs.slide_height)
