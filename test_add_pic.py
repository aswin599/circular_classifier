import pptx
from pptx.util import Inches

prs = pptx.Presentation('TCE_Linkedin_Project_Viva.pptx')
slide = prs.slides[0]
print("Testing add picture on dummy slide...")
test_slide = prs.slides.add_slide(prs.slide_layouts[6])
pic = test_slide.shapes.add_picture('screenshot_login.png', Inches(1), Inches(1), Inches(4), Inches(2.5))
print("Added picture:", pic.name, pic.width, pic.height)
