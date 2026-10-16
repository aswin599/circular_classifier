import pptx

prs = pptx.Presentation('TCE_Linkedin_Project_Viva.pptx')
for idx in [8, 9]:
    slide = prs.slides[idx]
    for shape in slide.shapes:
        if shape.has_chart:
            chart = shape.chart
            print(f"Slide {idx+1} Chart: title={chart.chart_title.text_frame.text if chart.has_title else 'No Title'}")
            print(f"Chart type: {chart.chart_type}")
            for series in chart.series:
                print(f"  Series: {series.name}, values={series.values}")
