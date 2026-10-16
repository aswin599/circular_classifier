import pptx

prs = pptx.Presentation('TCE_Linkedin_Project_Viva.pptx')
for idx in [0, 1, 4, 6, 8]:
    slide = prs.slides[idx]
    print(f"=== SLIDE {idx+1} ===")
    for shape in slide.shapes:
        if shape.has_text_frame:
            tf = shape.text_frame
            for p in tf.paragraphs[:2]:
                fn = p.font.name if p.font else "none"
                fs = p.font.size.pt if p.font and p.font.size else "none"
                runs_info = []
                for r in p.runs[:2]:
                    r_fn = r.font.name if r.font else "none"
                    r_fs = r.font.size.pt if r.font and r.font.size else "none"
                    r_c = str(r.font.color.rgb) if r.font and r.font.color and r.font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else "none"
                    runs_info.append(f"(f={r_fn}, sz={r_fs}, col={r_c})")
                print(f"  Shape {shape.name}: text='{p.text[:35]}', p_fn={fn}, runs={runs_info}")
