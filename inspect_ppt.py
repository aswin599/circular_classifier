import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = pptx.Presentation('TCE_Linkedin_Project_Viva.pptx')
print(f"Total slides: {len(prs.slides)}")
print(f"Slide width: {prs.slide_width}, height: {prs.slide_height}")

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for shape in slide.shapes:
        s_type = shape.shape_type
        pos = f"(x={shape.left/914400:.2f}\", y={shape.top/914400:.2f}\", w={shape.width/914400:.2f}\", h={shape.height/914400:.2f}\")"
        if shape.has_text_frame:
            full_text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            print(f"  Shape [{shape.name}] {pos}:\n{full_text}")
        elif shape.has_table:
            print(f"  Table [{shape.name}] {pos}: {len(shape.table.rows)} rows x {len(shape.table.columns)} cols")
        elif s_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"  Picture [{shape.name}] {pos}")
        elif s_type == MSO_SHAPE_TYPE.CHART:
            print(f"  Chart [{shape.name}] {pos}")
        else:
            print(f"  Other [{shape.name}] {pos} (type={s_type})")
