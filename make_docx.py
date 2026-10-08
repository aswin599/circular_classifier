# Rebuild the 44-page sample report as a Word document, keeping the original
# layout (margins, justification, centered headings, TOC tabs, tables, images,
# page numbers) while substituting the new project title and team members on
# pages 1-2.  Vertical placement is driven by the original PDF baselines:
#   space_before = prev_box_bottom - off(L, size) - desired_baseline
# with EXACT line spacing so Word reproduces the observed line rhythm.
import json
import os
import re

import content_edits as CE

import pypdfium2 as pdfium
from PIL import Image
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT,
                            WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = r"C:\Users\User\OneDrive\Desktop\Circular_Classification"
SRC = os.path.join(ROOT, "Project Report.pdf")
OUT = os.path.join(ROOT, "Circular Classifier - Project Report.docx")
STRUCT = json.load(open(os.path.join(ROOT, "_structure.json"), encoding="utf-8"))
PAGES = STRUCT["pages"]

TITLE = "CIRCULAR CLASSIFIER"
NAME1 = "THIRUNAVUKARASU S (2503917762221009)"
NAME2 = "ASWIN C (2503917762221001)"

# Replacement artwork for every image slot: drawn figures (exact slot aspect),
# app screenshots (top-cropped to slot aspect in fitted_image), crest/letter.
SWAP = {
    (1, 0): "_figs/logo.png",
    (2, 0): "_figs/logo.png",
    (3, 0): "_figs/declaration.png",
    (19, 0): "_figs/fig_arch.png",
    (20, 0): "_figs/fig_dfd0.png",
    (20, 1): "_figs/fig_dfd1.png",
    (21, 0): "_figs/fig_dfd2.png",
    (22, 0): "_figs/fig_component.png",
    (22, 1): "_figs/fig_usecase.png",
    (23, 0): "_figs/fig_deployment.png",
    # results section (5. results screenshots)
    (29, 0): "_figs/login.png",
    (29, 1): "_figs/register.png",
    (30, 0): "_figs/dashboard.png",
    (30, 1): "_figs/dashboard_ending.png",
    (30, 2): "_figs/detail.png",
    (31, 0): "_figs/detail2.png",
    (31, 1): "_figs/search.png",
    (31, 2): "_figs/admin_categories.png",
    (32, 0): "_figs/admin_users.png",
    (32, 1): "_figs/dashboard.png",
    (32, 2): "_figs/detail.png",
    (33, 0): "_figs/detail2.png",
    (33, 1): "_figs/search.png",
    (33, 2): "_figs/dashboard_ending.png",
    # appendix A - project screenshots
    (37, 0): "_figs/login.png",
    (37, 1): "_figs/register.png",
    (38, 0): "_figs/dashboard.png",
    (38, 1): "_figs/dashboard_ending.png",
    (38, 2): "_figs/detail.png",
    (38, 3): "_figs/detail2.png",
    (39, 0): "_figs/search.png",
    (39, 1): "_figs/admin_categories.png",
    (40, 0): "_figs/admin_users.png",
    (40, 1): "_figs/dashboard.png",
    (40, 2): "_figs/dashboard_ending.png",
    (41, 0): "_figs/detail.png",
    (41, 1): "_figs/detail2.png",
    (41, 2): "_figs/search.png",
    (42, 0): "_figs/admin_categories.png",
    (42, 1): "_figs/admin_users.png",
    (43, 0): "_figs/login.png",
    (43, 1): "_figs/register.png",
    (43, 2): "_figs/dashboard.png",
    (44, 0): "_figs/dashboard_ending.png",
    (44, 1): "_figs/detail.png",
    (44, 2): "_figs/detail2.png",
}

# Screenshot captions retitled to match the swapped images (left-aligned x0=90).
EDITS = [
    (29, 630.6, "LOGIN PAGE :"),
    (29, 358.9, "REGISTRATION :"),
    (30, 714.8, "DASHBOARD :"),
    (31, 726.7, "ATTACHMENT VIEW :"),
    (31, 483.4, "SEARCH & ADMIN :"),
    (32, 316.8, "CIRCULAR DETAILS :"),
    (32, 285.6, " "),
    (33, 350.8, "EXPIRY ALERTS :"),
    (37, 667.1, "LOGIN & DASHBOARD :"),
    (39, 740.7, "SEARCH & ADMIN :"),
    (40, 288.3, "EXPIRY, DETAIL & SEARCH :"),
    (42, 757.9, "ADMIN & LOGIN :"),
    (43, 555.9, "REGISTRATION :"),
    (43, 388.1, "DASHBOARD & DETAIL :"),
]
EDIT_MAP = {}
for _pg, _y, _t in EDITS:
    EDIT_MAP.setdefault(_pg, []).append((_y, _t))

PDF_H = 842.0
PAGE_W, PAGE_H = 595.28, 841.89          # Word A4
TOP, BOTTOM = 70.0, 36.0                 # margins (pt)
LEFT = 90.0
RIGHT = PAGE_W - 526.5                   # text right edge lands at 526.5
TEXT_RIGHT = 526.5
BODY_TOP = PDF_H - TOP

A, D, G = 0.891, 0.216, 0.033            # Times metrics (asc/desc/gap)

FONT_MAP = {
    "TimesNewRomanPS-BoldMT": ("Times New Roman", True, False),
    "TimesNewRomanPSMT": ("Times New Roman", False, False),
    "TimesNewRomanPS-ItalicMT": ("Times New Roman", False, True),
    "TimesNewRomanPS-BoldItalicMT": ("Times New Roman", True, True),
    "Cambria-BoldItalic": ("Cambria", True, True),
    "Cambria-Italic": ("Cambria", False, True),
    "Cambria-Bold": ("Cambria", True, False),
    "Calibri-Bold": ("Calibri", True, False),
    "Calibri": ("Calibri", False, False),
    "ArialMT": ("Arial", False, False),
    "Arial-BoldMT": ("Arial", True, False),
}

TOC_PAGES = {5, 6}
NUM_RE = re.compile(r"^(?P<t>.*\S)\s+(?P<n>[ivxlcdm]+|\d+)$", re.I)
LOG = []


def nat(s):
    return (A + D + G) * s


def off(L, s):
    # Word puts ~90% of the (L - nat) extra height ABOVE the first baseline
    # when exact line spacing L exceeds the font's natural line height
    # (verified against Word: para SpaceBefore exact, char vpos).
    return max((L - nat(s)) * 0.9, 0.0) + A * s


def font_of(name):
    return FONT_MAP.get(name, (name, False, False))


# ----------------------------------------------------------------- reflow
# Character advance widths (in em) measured directly from the source PDF, per
# font.  Used only to re-wrap rewritten paragraph text so that the number of
# rendered lines stays exactly the same as the original (geometry is anchored
# to the original baselines, so text must not change the line count).
def _build_widths():
    W = {}
    for p in PAGES:
        for x0, x1, y, f, s, _col, t in p["chars"]:
            if s and len(t) == 1 and t.strip():
                W.setdefault(f, {}).setdefault(t, []).append((x1 - x0) / s)
    for f in W:
        for ch in W[f]:
            v = W[f][ch]
            W[f][ch] = sum(v) / len(v)
    return W


WIDTHS = _build_widths()


def _adv(f, ch):
    tab = WIDTHS.get(f) or WIDTHS.get("TimesNewRomanPSMT", {})
    if ch in tab:
        return tab[ch]
    if ch == " ":
        return 0.25
    return tab.get("m", 0.6)


def _textw(s, f, size):
    return size * sum(_adv(f, c) for c in s)


def _dominant_font(lines):
    cnt = {}
    for l in lines:
        for r in l["runs"]:
            n = sum(1 for ch in r["t"] if not ch.isspace())
            if n:
                cnt[r["f"]] = cnt.get(r["f"], 0) + n
    return max(cnt, key=cnt.get) if cnt else "TimesNewRomanPSMT"


def _wrap(text, n, avails, f, size):
    words = text.split()
    if n <= 0:
        return None
    out, i = [], 0
    for k in range(n):
        rem = n - k
        avail = avails[k] if k < len(avails) else avails[-1]
        if i >= len(words):
            return None
        w, j = 0.0, i
        while j < len(words):
            add = _textw(words[j] if j == i else " " + words[j], f, size)
            if w + add > avail and j > i:
                break
            w += add
            j += 1
        maxj = len(words) - (rem - 1)
        if j > maxj:
            j = maxj
        if j <= i:
            j = i + 1
        out.append(" ".join(words[i:j]))
        i = j
    if i != len(words):
        return None
    return out


def apply_line_edits(lines, pno):
    for ln in lines:
        for (pg, y), t in CE.LINE_EDITS.items():
            if pg == pno and abs(ln["y"] - y) < 0.06:
                f = _dominant_font([ln])
                s = maxsize(ln)
                ln["runs"] = [{"f": f, "s": s, "t": t}]
                ln["text"] = t
                break
    return lines


def _block_lines(block):
    return [[l for l in pa["lines"] if l["text"].strip()] for pa in block]


def apply_block_edit(block, text, pno):
    allpl = _block_lines(block)
    flat = [l for pl in allpl for l in pl]
    if not flat:
        return
    f = _dominant_font(flat)
    size = maxsize(flat[0])
    avails = [TEXT_RIGHT - l["x0"] - 1.0 for l in flat]
    out = _wrap(text, len(flat), avails, f, size)
    if out is None:
        LOG.append("WRAP FAIL p%d y=%.1f n=%d" % (pno, flat[0]["y"], len(flat)))
        return
    k = 0
    for pl in allpl:
        for l in pl:
            l["runs"] = [{"f": f, "s": size, "t": out[k]}]
            l["text"] = out[k]
            k += 1


# ----------------------------------------------------------------- lines
def group_runs(chars):
    runs = []
    for x0, x1, y, font, size, color, text in chars:
        if runs and runs[-1]["f"] == font and runs[-1]["s"] == size:
            runs[-1]["t"] += text
        else:
            runs.append({"f": font, "s": size, "t": text})
    return runs


def page_lines(pno, skip_y=()):
    p = PAGES[pno - 1]
    buckets = {}
    for c in p["chars"]:
        buckets.setdefault(c[2], []).append(c)
    lines = []
    for y in sorted(buckets, reverse=True):
        if y in skip_y:
            continue
        cs = sorted(buckets[y], key=lambda c: c[0])
        if any(t["y_bot"] - 0.1 <= y <= t["y_top"] + 0.1
               and cs[0][0] >= t["x0"] - 3 and cs[-1][1] <= t["x1"] + 3
               for t in p["tables"]):
            continue
        runs = group_runs(cs)
        ink = [c for c in cs if c[6].strip()]
        lines.append({
            "y": y,
            "x0": cs[0][0],
            "x1": max(c[1] for c in ink) if ink else cs[-1][1],
            "runs": runs,
            "text": "".join(r["t"] for r in runs),
        })
    return lines


def maxsize(line):
    return max(r["s"] for r in line["runs"])


def bold_ratio(line):
    nb = bb = 0
    for r in line["runs"]:
        for ch in r["t"]:
            if not ch.isspace():
                nb += 1
                if "Bo" in r["f"]:
                    bb += 1
    return bb / nb if nb else 0.0


def is_heading(line):
    # A fully-bold line is a section heading. Headings must stay solitary
    # paragraphs: joining one to the body text lets Word flow heading+body
    # as continuous wrapped text, saving one visual line per section and
    # dragging everything below out of position (up to ~100pt per page).
    if "_h" not in line:
        line["_h"] = bold_ratio(line) >= 0.8
    return line["_h"]


def group_paras(lines, pno):
    paras, cur = [], None
    for ln in lines:
        join = False
        if cur and pno not in TOC_PAGES:
            prev = cur["lines"][-1]
            gap = prev["y"] - ln["y"]
            # Whitespace-only baselines are blank lines in the original:
            # they are hard paragraph boundaries (joining across them
            # corrupts L and the justification heuristic), and they render
            # as pure gap (S = G - L >= 0 on the uniform line grid).
            join = (gap > 0 and gap <= 1.95 * maxsize(prev)
                    and maxsize(ln) == maxsize(prev)
                    and ln["x0"] >= prev["x0"]
                    and ln["x0"] - min(l["x0"] for l in cur["lines"]) <= 30
                    and prev["text"].strip() and ln["text"].strip()
                    and not is_heading(prev) and not is_heading(ln))
        if join:
            cur["lines"].append(ln)
        else:
            if cur:
                paras.append(cur)
            cur = {"lines": [ln]}
            for k in ("L", "render_n", "force_justify"):
                if k in ln:
                    cur[k] = ln[k]
    if cur:
        paras.append(cur)
    return paras


def para_items(lines, breaks=True):
    """[(run_dict_or_None for separator)] -> merged [(style, text)] items."""
    items = []
    for i, ln in enumerate(lines):
        if i > 0:
            if breaks:
                # Force Word to wrap exactly where the source PDF wrapped:
                # natural re-wrapping can shuffle boundary words and drop a
                # line (or add one), shifting everything below by ~L.
                items.append({"_br": True})
            elif items:
                items[-1]["t"] += " "
        rs = [dict(r) for r in ln["runs"]]
        if rs:
            rs[-1]["t"] = rs[-1]["t"].rstrip()
        if i > 0 and rs:
            rs[0]["t"] = rs[0]["t"].lstrip()
        for r in rs:
            if not r["t"]:
                continue
            if (items and "_br" not in items[-1]
                    and items[-1]["f"] == r["f"] and items[-1]["s"] == r["s"]):
                items[-1]["t"] += r["t"]
            else:
                items.append(r)
    return items


# ----------------------------------------------------------------- docx bits
def set_spacing(pf, L, S, indent, first_indent, justify, page_break,
                rindent=0.0):
    pf.space_before = Pt(max(S, 0.0))
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(L)
    pf.line_spacing_rule = 4  # EXACTLY (wdLineSpacingExactly)
    pf.left_indent = Pt(indent)
    pf.right_indent = Pt(rindent)
    pf.first_line_indent = Pt(first_indent)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    pf.page_break_before = bool(page_break)
    pPr = pf._element.get_or_add_pPr()
    for tag in ("w:widowControl", "w:snapToGrid"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "0")
        pPr.append(el)


def add_runs(p, items, color=None):
    for r in items:
        if r.get("_br"):
            p.add_run()._r.append(OxmlElement("w:br"))
            continue
        fam, bold, italic = font_of(r["f"])
        run = p.add_run(r["t"])
        run.font.name = fam
        run.font.size = Pt(r["s"])
        run.font.bold = bold
        run.font.italic = italic
        if color is not None:
            run.font.color.rgb = color
    if not items:
        p.add_run("")


def fill_hex(fill):
    if not fill:
        return None
    return "%02X%02X%02X" % tuple(int(round(v * 255)) for v in fill[:3])


def lum(fill):
    return 0.299 * fill[0] + 0.587 * fill[1] + 0.114 * fill[2]


def cell_border(color):
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), color)
        b.append(e)
    return b


def cell_shade(hexcolor):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    return sh


def cell_margins_zero(tc, top=0.0):
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("start", 0), ("bottom", 0), ("end", 0)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(int(round(val * 20))))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


# ----------------------------------------------------------------- images
_pdf = pdfium.PdfDocument(SRC)
_rend = {}
IMG_DIR = os.path.join(ROOT, "_docx_imgs")
os.makedirs(IMG_DIR, exist_ok=True)


def page_image(pno):
    if pno not in _rend:
        _rend[pno] = _pdf[pno - 1].render(scale=3).to_pil().convert("RGB")
    return _rend[pno]


def image_png(pno, idx, im):
    key = (pno, idx)
    if key in SWAP:
        src = os.path.join(ROOT, SWAP[key])
        out = os.path.join(IMG_DIR, "swap_p%02d_%d.png" % (pno, idx))
        return fitted_image(src, im["x1"] - im["x0"], im["y_top"] - im["y_bot"],
                            out)
    bmp = page_image(pno)
    W, H = bmp.size
    sx, sy = W / 596.0, H / 842.0
    box = (round(im["x0"] * sx), round((PDF_H - im["y_top"]) * sy),
           round(im["x1"] * sx), round((PDF_H - im["y_bot"]) * sy))
    path = os.path.join(IMG_DIR, "p%02d_%d.png" % (pno, idx))
    bmp.crop(box).save(path)
    return path


def fitted_image(src, slot_w, slot_h, out):
    """Reproduce the slot aspect exactly so the layout math stays valid.
    Matching artwork passes through untouched; wider-than-source slots are
    top-cropped (keeps browser/app header), taller ones get white padding."""
    im = Image.open(src).convert("RGB")
    A = slot_w / slot_h
    sw, sh = im.size
    if abs(sw / sh - A) <= 0.01 * A:
        im.save(out)
        return out
    new_h = int(round(sw / A))
    if new_h <= sh:
        im = im.crop((0, 0, sw, new_h))
    else:
        canvas = Image.new("RGB", (sw, new_h), "white")
        canvas.paste(im, (0, 0))
        im = canvas
    im.save(out)
    return out


# ----------------------------------------------------------------- table
def add_table(doc, t, page_chars, pno):
    rows = t["rows"]
    nR, nC = len(rows), len(rows[0]["cells"])
    tbl = doc.add_table(nR, nC)
    tbl.autofit = False
    widths = [c[1] - c[0] for c in rows[0]["cells"]]
    has_fill = any(r["fill"] for r in rows)
    if not has_fill:
        tbl.style = "Table Grid"
    for ri, row in enumerate(rows):
        rx0, rx1, rtop, rbot = row["bbox"]
        r = tbl.rows[ri]
        r.height = Pt(rtop - rbot)
        r.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        fill = row["fill"]
        hexf = fill_hex(fill)
        white_text = bool(fill) and lum(fill) < 0.6
        for ci, cbox in enumerate(row["cells"]):
            cell = tbl.cell(ri, ci)
            cell.width = Pt(widths[ci])
            tcPr = cell._tc.get_or_add_tcPr()
            if hexf:
                tcPr.append(cell_shade(hexf))
                tcPr.append(cell_border(hexf))
            if cbox is None:
                cell_margins_zero(cell._tc)
                continue
            cx0, cx1, ctop, cbot = cbox
            cc = [c for c in page_chars
                  if cbot - 0.1 <= c[2] <= ctop + 0.1
                  and c[0] >= cx0 - 0.6 and c[1] <= cx1 + 0.6]
            if not cc:
                cell_margins_zero(cell._tc)
                continue
            buckets = {}
            for c in cc:
                buckets.setdefault(c[2], []).append(c)
            clines = []
            for y in sorted(buckets, reverse=True):
                cs = sorted(buckets[y], key=lambda c: c[0])
                runs = group_runs(cs)
                ink = [c for c in cs if c[6].strip()]
                clines.append({
                    "y": y, "x0": cs[0][0],
                    "x1": max(c[1] for c in ink) if ink else cs[-1][1],
                    "runs": runs,
                    "text": "".join(r["t"] for r in runs),
                })
            tkey = (pno, ri, ci)
            if tkey in CE.TABLE_EDITS:
                tf = _dominant_font(clines)
                tsz = maxsize(clines[0])
                tL = (clines[0]["y"] - clines[-1]["y"]) / (len(clines) - 1) \
                    if len(clines) >= 2 else 1.15 * nat(tsz)
                avail = cx1 - cx0 - 2.0
                wl, cur = [], ""
                for w in CE.TABLE_EDITS[tkey].split():
                    t = (cur + " " + w).strip()
                    if not cur or _textw(t, tf, tsz) <= avail:
                        cur = t
                    else:
                        wl.append(cur)
                        cur = w
                if cur:
                    wl.append(cur)
                if len(wl) > len(clines):
                    LOG.append("TABLE OVERFLOW p%d r%d c%d %d>%d"
                               % (pno, ri, ci, len(wl), len(clines)))
                    wl = wl[:len(clines)]
                ncl, yy = [], clines[0]["y"]
                for txt in wl:
                    ncl.append({"y": yy, "x0": clines[0]["x0"],
                                "x1": clines[0]["x0"] + _textw(txt, tf, tsz),
                                "runs": [{"f": tf, "s": tsz, "t": txt}],
                                "text": txt})
                    yy -= tL
                clines = ncl
            p = cell.paragraphs[0]
            s = maxsize(clines[0])
            if len(clines) >= 2:
                L = (clines[0]["y"] - clines[-1]["y"]) / (len(clines) - 1)
            else:
                L = 1.15 * nat(s)
            S = rtop - off(L, s) - clines[0]["y"]
            body_x0 = (clines[-1]["x0"] if len(clines) > 1 else clines[0]["x0"])
            # Cells are written from the top with space-before equal to the
            # source top offset (each wrapped line is a hard line break, so
            # the emitted line count is exact).
            cell_margins_zero(cell._tc)
            set_spacing(p.paragraph_format, L, S, body_x0 - cx0,
                        clines[0]["x0"] - body_x0, False, False)
            add_runs(p, para_items(clines),
                     color=RGBColor(255, 255, 255) if white_text else None)
    return tbl


# ----------------------------------------------------------------- page 2
def page2_runs():
    R = [
        ("TimesNewRomanPSMT", 14, "Certified that this project report \u201c"),
        ("TimesNewRomanPS-BoldMT", 14, TITLE + "\u201d"),
        ("TimesNewRomanPSMT", 14, " is the bonafide work of \u201c"),
        ("TimesNewRomanPS-BoldMT", 16, NAME1 + ","),
        ("TimesNewRomanPSMT", 14, " "),
        ("TimesNewRomanPS-BoldMT", 16, NAME2),
        ("TimesNewRomanPS-BoldMT", 14, "\u201d"),
        ("TimesNewRomanPSMT", 14,
         " who carried out the Mini project work under my supervision "
         "during the Academic Year 2025-2026."),
    ]
    return [{"f": f, "s": s, "t": t} for f, s, t in R]


def page1_edits(lines):
    out = []
    for ln in lines:
        if abs(ln["y"] - 722.1) < 0.05:          # WEBSITE DEVELOPMENT removed
            continue
        if abs(ln["y"] - 753.2) < 0.05:
            ln = dict(ln, x0=192.9, x1=402.4,
                      runs=[{"f": "TimesNewRomanPS-BoldMT", "s": 18.0,
                             "t": TITLE}], text=TITLE)
        elif abs(ln["y"] - 593.0) < 0.05:
            ln = dict(ln, x0=150.8, x1=472.6,
                      runs=[{"f": "TimesNewRomanPS-BoldMT", "s": 16.0,
                             "t": NAME1}], text=NAME1)
        elif abs(ln["y"] - 574.6) < 0.05:
            ln = dict(ln, x0=203.8, x1=416.2,
                      runs=[{"f": "TimesNewRomanPS-BoldMT", "s": 16.0,
                             "t": NAME2}], text=NAME2)
        out.append(ln)
    return out


def toc_parts(line):
    text = line["text"].rstrip()
    m = NUM_RE.match(text)
    if not m:
        return None
    num = m.group("n")
    if len(num) > 4:
        return None
    t_txt = m.group("t")
    num_start = m.start("n")
    # split runs at num_start
    left, right, pos = [], [], 0
    for r in line["runs"]:
        end = pos + len(r["t"])
        if end <= num_start:
            left.append(dict(r))
        elif pos >= num_start:
            right.append(dict(r))
        else:
            a, b = num_start - pos, num_start - pos
            left.append({**r, "t": r["t"][:a]})
            right.append({**r, "t": r["t"][b:]})
        pos = end
    while left and not left[-1]["t"].strip():
        left.pop()
    while left and left[-1]["t"].endswith(" "):
        left[-1]["t"] = left[-1]["t"].rstrip()
    right = [r for r in right if r["t"]]
    return left, right, line["x1"] - LEFT


# ----------------------------------------------------------------- main
def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(0)
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Pt(TOP)
    sec.bottom_margin = Pt(BOTTOM)
    sec.left_margin = Pt(LEFT)
    sec.right_margin = Pt(RIGHT)

    neg = 0
    for pno in range(1, 45):
        p = PAGES[pno - 1]
        if pno == 1:
            lines = page_lines(1)
            lines = page1_edits(lines)
        elif pno == 2:
            lines = page_lines(2, skip_y=(552.3, 526.3, 498.7, 473.0, 448.8))
            lines.insert(0, {
                "y": 552.3, "x0": 108.0, "x1": TEXT_RIGHT,
                "runs": page2_runs(),
                "text": " ".join(r["t"] for r in page2_runs()),
                "L": (552.3 - 448.8) / 4.0,
                "render_n": 5,
                "force_justify": True,
            })
        else:
            lines = page_lines(pno)

        lines = apply_line_edits(lines, pno)

        for ln in lines:
            for y, t in EDIT_MAP.get(pno, ()):
                if abs(ln["y"] - y) < 0.06:
                    bold = [r for r in ln["runs"]
                            if r["f"] == "TimesNewRomanPS-BoldMT"]
                    r = bold[-1] if bold else ln["runs"][-1]
                    ln["runs"] = [{"f": r["f"], "s": r["s"], "t": t}]
                    ln["text"] = t
                    break

        paras = group_paras(lines, pno)
        i = 0
        while i < len(paras):
            pa = paras[i]
            pl = [l for l in pa["lines"] if l["text"].strip()]
            if not pl:
                i += 1
                continue
            key = (pno, round(pl[0]["y"], 1))
            if key not in CE.PARA_EDITS:
                i += 1
                continue
            text = CE.PARA_EDITS[key]
            # An indented first line and its continuation are stored as two
            # separate source paragraphs; a rewritten full paragraph therefore
            # has to span them. Absorb following paragraphs (contiguous, same
            # size, non-heading) only until the whole text fits exactly.
            block = [pa]
            j = i + 1
            while True:
                allpl = _block_lines(block)
                flat = [l for pl in allpl for l in pl]
                f = _dominant_font(flat)
                size = maxsize(flat[0])
                avails = [TEXT_RIGHT - l["x0"] - 1.0 for l in flat]
                if _wrap(text, len(flat), avails, f, size) is not None:
                    break
                if j >= len(paras):
                    break
                nxt = paras[j]
                npl = [l for l in nxt["lines"] if l["text"].strip()]
                if not npl:
                    break
                last = [l for l in block[-1]["lines"] if l["text"].strip()][-1]
                gap = last["y"] - npl[0]["y"]
                if (0 < gap <= 1.95 * maxsize(last)
                        and maxsize(npl[0]) == maxsize(last)
                        and not is_heading(npl[0]) and not is_heading(last)
                        and (pno, round(npl[0]["y"], 1)) not in CE.PARA_EDITS):
                    block.append(nxt)
                    j += 1
                else:
                    break
            apply_block_edit(block, text, pno)
            i = j if j > i + 1 else i + 1
        tables = p["tables"]
        imgs = p["images"]
        elements = []
        for pa in paras:
            elements.append((pa["lines"][0]["y"], "p", pa))
        for i, im in enumerate(imgs):
            elements.append((im["y_top"], "i", (i, im)))
        for t in tables:
            elements.append((t["y_top"], "t", t))
        elements.sort(key=lambda e: -e[0])

        Yb = BODY_TOP
        first = True
        last_p = None
        for _, kind, obj in elements:
            if kind == "p":
                plines = obj["lines"]
                # Whitespace-only baselines (spacer rows in the source
                # grid) contribute no rendered line: keep only text lines
                # so L/rn match what Word actually draws. Keeping a
                # trailing ws row would book a box one line too tall and
                # shift everything below up by ~L per occurrence.
                plines = [l for l in plines if l["text"].strip()]
                if not plines:
                    continue
                L = obj.get("L")
                if L is None:
                    if len(plines) >= 2:
                        L = (plines[0]["y"] - plines[-1]["y"]) / (len(plines) - 1)
                    else:
                        L = 1.15 * maxsize(plines[0])
                s0 = maxsize(plines[0])
                B = plines[0]["y"]
                S = Yb - off(L, s0) - B
                if S < -0.5:
                    neg += 1
                    LOG.append("p%d neg S %.1f (B=%.1f)" % (pno, S, B))
                items = para_items(plines)
                if pno in TOC_PAGES:
                    sp = toc_parts(plines[0])
                    pp = doc.add_paragraph()
                    pf = pp.paragraph_format
                    if sp:
                        left, right, tabpos = sp
                        set_spacing(pf, L, S, plines[0]["x0"] - LEFT, 0,
                                    False, pno > 1 and first)
                        add_runs(pp, left)
                        tabr = pp.add_run()
                        tabr._r.append(OxmlElement("w:tab"))
                        pf.tab_stops.add_tab_stop(Pt(tabpos),
                                                  WD_TAB_ALIGNMENT.RIGHT,
                                                  WD_TAB_LEADER.SPACES)
                        add_runs(pp, right)
                    else:
                        body_x0 = plines[0]["x0"]
                        set_spacing(pf, L, S, body_x0 - LEFT, 0, False,
                                    pno > 1 and first)
                        add_runs(pp, items)
                else:
                    n = len(plines)
                    body_x0 = (plines[-1]["x0"] if n > 1 else plines[0]["x0"])
                    justify = obj.get("force_justify") or (
                        n >= 2 and all(l["x1"] >= 518 for l in plines[:-1]))
                    max_x1 = max(l["x1"] for l in plines)
                    rindent = min(0.0, TEXT_RIGHT - max_x1 - 1.0)
                    pp = doc.add_paragraph()
                    set_spacing(pp.paragraph_format, L, S,
                                body_x0 - LEFT, plines[0]["x0"] - body_x0,
                                justify, pno > 1 and first, rindent)
                    add_runs(pp, items)
                rn = obj.get("render_n", len(plines))
                Yb = B + off(L, s0) - rn * L
                last_p = pp
            elif kind == "i":
                idx, im = obj
                path = image_png(pno, idx, im)
                # Word places a picture's baseline at 0.8*L below the
                # paragraph top (verified empirically for EXACT line
                # spacing), with the picture bottom on that baseline.
                # Use a tiny line box L: picture top lands at
                # y_top = (y_bot + 0.8L) - 0.8L + h, the box bottom stays
                # 0.2L below y_bot, and the picture simply overflows the
                # box upward - keeping following content unclamped.
                IL = 10.0
                S = Yb - im["y_bot"] - 0.8 * IL
                if S < -0.5:
                    neg += 1
                    LOG.append("p%d img neg S %.1f" % (pno, S))
                pp = doc.add_paragraph()
                pf = pp.paragraph_format
                pf.space_before = Pt(max(S, 0.0))
                pf.space_after = Pt(0)
                pf.line_spacing = Pt(IL)
                pf.line_spacing_rule = 4
                pf.left_indent = Pt(im["x0"] - LEFT)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.page_break_before = bool(pno > 1 and first)
                pPr = pf._element.get_or_add_pPr()
                for tag in ("w:widowControl", "w:snapToGrid"):
                    el = OxmlElement(tag)
                    el.set(qn("w:val"), "0")
                    pPr.append(el)
                pp.add_run().add_picture(path, width=Pt(im["x1"] - im["x0"]))
                Yb = im["y_bot"] - 0.2 * IL
                last_p = pp
            else:  # table
                S = Yb - obj["y_top"]
                if S < -0.5:
                    neg += 1
                    LOG.append("p%d table neg S %.1f" % (pno, S))
                if last_p is not None:
                    last_p.paragraph_format.space_after = Pt(max(S, 0.0))
                add_table(doc, obj, p["chars"], pno)
                Yb = obj["y_bot"]
                last_p = None
            first = False

    doc.save(OUT)
    print("wrote", OUT)
    if neg:
        print("negative spacings:", neg)
    if LOG:
        print("issues:", len(LOG))
        for line in LOG[:80]:
            print("  ", line)


if __name__ == "__main__":
    main()
